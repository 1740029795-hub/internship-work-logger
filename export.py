"""
导出模块 — Word (.docx) 和 PDF
"""
import io
import os
from datetime import datetime

from database import get_db, get_settings


def _get_report_data(report_type, report_id):
    """获取报告数据"""
    conn = get_db()
    row = None
    if report_type == 'weekly':
        row = conn.execute('SELECT * FROM weekly_reports WHERE id = ?', (report_id,)).fetchone()
    elif report_type == 'monthly':
        row = conn.execute('SELECT * FROM monthly_reports WHERE id = ?', (report_id,)).fetchone()
    elif report_type == 'annual':
        row = conn.execute('SELECT * FROM annual_reports WHERE id = ?', (report_id,)).fetchone()
    conn.close()

    if row:
        return dict(row)
    return None


def export_to_docx(report_type, report_id):
    """
    导出 Word (.docx) 文件
    返回 BytesIO 对象
    """
    from docx import Document
    from docx.shared import Pt, Inches, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT

    report = _get_report_data(report_type, report_id)
    settings = get_settings()

    doc = Document()

    # 页面设置
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)

    # 样式设置
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Microsoft YaHei'
    font.size = Pt(11)
    style.paragraph_format.line_spacing = 1.5

    # 标题
    type_labels = {'weekly': '实习周报', 'monthly': '实习月报', 'annual': '实习年报'}
    title_text = type_labels.get(report_type, '实习报告')

    title = doc.add_heading(title_text, level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.size = Pt(22)
        run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    # 副标题：日期范围
    subtitle_text = ''
    if report_type == 'weekly' and report:
        subtitle_text = f"{report['week_start']} 至 {report['week_end']}  |  第 {report['week_number']} 周"
    elif report_type == 'monthly' and report:
        subtitle_text = f"{report['report_year']}年{report['report_month']}月"
    elif report_type == 'annual' and report:
        subtitle_text = f"{report['report_year']}年"

    if subtitle_text:
        subtitle = doc.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = subtitle.add_run(subtitle_text)
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    # 空行
    doc.add_paragraph()

    # 学生信息表
    if settings:
        table = doc.add_table(rows=4, cols=4)
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        info_pairs = [
            ('姓名', settings.get('student_name', '')),
            ('学号', settings.get('student_id', '')),
            ('学校', settings.get('school', '')),
            ('专业', settings.get('major', '')),
            ('实习单位', settings.get('company_name', '')),
            ('岗位', settings.get('position', '')),
            ('部门', settings.get('department', '')),
            ('指导老师', settings.get('mentor_name', '')),
        ]

        for i, (label, value) in enumerate(info_pairs):
            row_idx = i // 2
            col_idx = (i % 2) * 2
            cell_label = table.cell(row_idx, col_idx)
            cell_value = table.cell(row_idx, col_idx + 1)

            cell_label.text = label
            cell_value.text = value or '（未填写）'

            # 样式
            for paragraph in cell_label.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    run.font.bold = True
            for paragraph in cell_value.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)

        # 设置列宽
        for row in table.rows:
            row.cells[0].width = Cm(3)
            row.cells[1].width = Cm(5.5)
            row.cells[2].width = Cm(3)
            row.cells[3].width = Cm(5.5)

    # 分隔线
    doc.add_paragraph('—' * 40)

    # 报告内容
    if report:
        section_defs = {
            'weekly': [
                ('本周摘要', 'summary'),
                ('工作亮点', 'work_highlights'),
                ('下周计划', 'next_week_plan'),
            ],
            'monthly': [
                ('月度总结', 'summary'),
                ('主要成果', 'achievements'),
                ('技能收获', 'skills_gained'),
                ('反思与不足', 'reflection'),
            ],
            'annual': [
                ('年度总结', 'summary'),
                ('总体回顾', 'overall_review'),
                ('重点项目', 'key_projects'),
                ('成长总结', 'growth_summary'),
                ('未来展望', 'future_outlook'),
            ],
        }

        for section_title, key in section_defs.get(report_type, []):
            h = doc.add_heading(section_title, level=2)
            for run in h.runs:
                run.font.size = Pt(14)
                run.font.color.rgb = RGBColor(0x2C, 0x52, 0x82)

            content = report.get(key, '') or '（无内容）'
            p = doc.add_paragraph(content)
            p.paragraph_format.line_spacing = 1.8

    # 页脚
    doc.add_paragraph()
    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(f"生成日期：{datetime.now().strftime('%Y-%m-%d %H:%M')}  |  实习工作记录器")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def render_report_html(report_type, report_id):
    """
    渲染报告的 HTML（用于 PDF 导出和打印回退）
    返回 HTML 字符串
    """
    report = _get_report_data(report_type, report_id)
    settings = get_settings()

    type_labels = {'weekly': '实习周报', 'monthly': '实习月报', 'annual': '实习年报'}
    title_text = type_labels.get(report_type, '实习报告')

    # 日期范围和标签
    range_text = ''
    if report_type == 'weekly' and report:
        range_text = f'{report["week_start"]} 至 {report["week_end"]} ｜ 第 {report["week_number"]} 周'
    elif report_type == 'monthly' and report:
        range_text = f'{report["report_year"]}年{report["report_month"]}月'
    elif report_type == 'annual' and report:
        range_text = f'{report["report_year"]}年'

    section_defs = {
        'weekly': [
            ('本周摘要', 'summary'),
            ('工作亮点', 'work_highlights'),
            ('下周计划', 'next_week_plan'),
        ],
        'monthly': [
            ('月度总结', 'summary'),
            ('主要成果', 'achievements'),
            ('技能收获', 'skills_gained'),
            ('反思与不足', 'reflection'),
        ],
        'annual': [
            ('年度总结', 'summary'),
            ('总体回顾', 'overall_review'),
            ('重点项目', 'key_projects'),
            ('成长总结', 'growth_summary'),
            ('未来展望', 'future_outlook'),
        ],
    }

    # 构建内容部分 HTML
    sections_html = ''
    if report:
        for section_title, key in section_defs.get(report_type, []):
            content = (report.get(key, '') or '（无内容）').replace('\n', '<br>')
            sections_html += f'''
            <h2>{section_title}</h2>
            <div class="section-content">{content}</div>
            '''

    html = f'''<!DOCTYPE html>
<html lang="zh-CN">
<head>
    <meta charset="UTF-8">
    <title>{title_text} - {settings.get("student_name", "学生")}</title>
    <style>
        @page {{
            size: A4;
            margin: 2cm;
            @bottom-center {{
                content: "生成日期：{datetime.now().strftime("%Y-%m-%d")} ｜ 实习工作记录器";
                font-size: 9pt;
                color: #999;
            }}
        }}
        body {{
            font-family: "Microsoft YaHei", "SimHei", "PingFang SC", sans-serif;
            font-size: 12pt;
            line-height: 1.8;
            color: #333;
        }}
        h1 {{
            font-size: 22pt;
            text-align: center;
            color: #1a365d;
            margin-bottom: 0;
        }}
        .subtitle {{
            text-align: center;
            color: #666;
            font-size: 11pt;
            margin-bottom: 24px;
        }}
        .info-table {{
            width: 100%;
            border-collapse: collapse;
            margin: 16px 0 24px;
        }}
        .info-table td {{
            padding: 6px 10px;
            border: 1px solid #ddd;
            font-size: 10pt;
        }}
        .info-table .label {{
            background: #f7fafc;
            font-weight: bold;
            width: 15%;
        }}
        h2 {{
            font-size: 15pt;
            color: #2c5282;
            border-bottom: 2px solid #2c5282;
            padding-bottom: 4px;
            margin-top: 24px;
        }}
        .section-content {{
            margin: 8px 0 16px;
            padding-left: 8px;
        }}
        .footer {{
            margin-top: 40px;
            text-align: center;
            font-size: 9pt;
            color: #999;
            border-top: 1px solid #eee;
            padding-top: 10px;
        }}
    </style>
</head>
<body>
    <h1>{title_text}</h1>
    <p class="subtitle">{range_text}</p>

    <table class="info-table">
        <tr>
            <td class="label">姓名</td><td>{settings.get('student_name', '')}</td>
            <td class="label">学号</td><td>{settings.get('student_id', '')}</td>
        </tr>
        <tr>
            <td class="label">学校</td><td>{settings.get('school', '')}</td>
            <td class="label">专业</td><td>{settings.get('major', '')}</td>
        </tr>
        <tr>
            <td class="label">实习单位</td><td>{settings.get('company_name', '')}</td>
            <td class="label">岗位</td><td>{settings.get('position', '')}</td>
        </tr>
        <tr>
            <td class="label">部门</td><td>{settings.get('department', '')}</td>
            <td class="label">指导老师</td><td>{settings.get('mentor_name', '')}</td>
        </tr>
    </table>

    <hr>

    {sections_html}

    <div class="footer">
        生成日期：{datetime.now().strftime('%Y-%m-%d %H:%M')} ｜ 实习工作记录器
    </div>
</body>
</html>'''

    return html


def export_to_pdf(report_type, report_id):
    """
    导出 PDF 文件
    使用 WeasyPrint 渲染 HTML → PDF
    返回 bytes
    """
    from weasyprint import HTML

    html_content = render_report_html(report_type, report_id)
    pdf_bytes = HTML(string=html_content).write_pdf()
    return pdf_bytes
